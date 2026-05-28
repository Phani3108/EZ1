"""Reports — read-only API (PH2-11).

This router replaces the pre-consolidation reporting-service HTTP surface.
Endpoints kept (same paths, same shapes — public contract is unchanged):

  GET  /reports/dashboard
  GET  /reports/attendance/trend
  GET  /reports/financial/summary
  GET  /reports/dropout/summary
  GET  /reports/dropout/students
  GET  /reports/dropout/student/{student_id}
  GET  /reports/export/attendance         (Excel)
  GET  /reports/export/financial          (Excel)
  GET  /reports/export/dropout            (Excel)
  GET  /reports/export/report-card/{id}   (PDF)

Endpoints DROPPED (now CLI-only — see services/reporting-service/cli/):

  POST /reports/consume   → cli/consume_events.py
  POST /reports/rebuild   → cli/rebuild_projections.py

The dashboard / trend / financial reads run against the reporting
projection DB via a second SQLAlchemy connection (`reporting_db.py`). The
dropout endpoints read students + attendance in-process (academics owns
those tables) and call finance over HTTP for invoice data.
"""
from __future__ import annotations

import io
import logging
import uuid
from datetime import date, datetime, timedelta, timezone
from typing import Optional

import httpx
from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.config import get_settings
from app.database import get_db
from app.dependencies import get_current_user, get_school_id
# Phase 20a — shared route helpers.
from eduzim_shared.routes import (
    _meta,
)
from app.reporting_db import get_reporting_db
from app.services.reports.dropout_data import (
    LOOKBACK_DAYS,
    compute_student_risk,
    fetch_invoices,
    fetch_student_trend,
    list_active_students,
)
from app.services.reports.dropout_engine import DropoutRiskEngine
from app.services.reports.reporting_query_service import ReportingQueryService


logger = logging.getLogger(__name__)

router = APIRouter(tags=["Reports"])

settings = get_settings()



def _extract_token(request: Request) -> str:
    auth = request.headers.get("authorization", "")
    if auth.startswith("Bearer "):
        return auth[7:]
    return ""


# ───────────────────── Dashboard / Trend / Financial ───────────────────────


@router.get("/reports/dashboard")
def dashboard(
    request: Request,
    rdb: Session = Depends(get_reporting_db),
    school_id: uuid.UUID = Depends(get_school_id),
):
    svc = ReportingQueryService(rdb)
    return {"data": svc.get_dashboard(school_id), "meta": _meta(request)}


@router.get("/reports/attendance/trend")
def attendance_trend(
    request: Request,
    from_date: date = Query(..., alias="from"),
    to_date: date = Query(..., alias="to"),
    rdb: Session = Depends(get_reporting_db),
    school_id: uuid.UUID = Depends(get_school_id),
):
    svc = ReportingQueryService(rdb)
    return {
        "data": svc.get_attendance_trend(school_id, from_date, to_date),
        "meta": _meta(request),
    }


@router.get("/reports/financial/summary")
def financial_summary(
    request: Request,
    year_id: Optional[uuid.UUID] = Query(None),
    rdb: Session = Depends(get_reporting_db),
    school_id: uuid.UUID = Depends(get_school_id),
):
    svc = ReportingQueryService(rdb)
    return {
        "data": svc.get_financial_summary(school_id, year_id),
        "meta": _meta(request),
    }


# ───────────────────── Dropout Intelligence ───────────────────────


@router.get("/reports/dropout/summary")
async def dropout_summary(
    request: Request,
    db: Session = Depends(get_db),
    school_id: uuid.UUID = Depends(get_school_id),
):
    """Aggregated dropout risk summary for the school."""
    token = _extract_token(request)
    today = date.today()

    students = list_active_students(db, school_id)
    band_counts = {"LOW": 0, "MEDIUM": 0, "HIGH": 0, "CRITICAL": 0}
    signal_freq: dict[str, int] = {}
    at_risk = 0

    for student in students:
        sid_str = student.get("id")
        if not sid_str:
            continue
        try:
            sid_uuid = uuid.UUID(sid_str)
        except (TypeError, ValueError):
            continue
        risk = await compute_student_risk(db, school_id, sid_uuid, token, today)
        band = risk["risk_band"]
        band_counts[band] = band_counts.get(band, 0) + 1
        if band != "LOW":
            at_risk += 1
        for sig in risk["signals"]:
            code = sig["code"]
            signal_freq[code] = signal_freq.get(code, 0) + 1

    top_signals = sorted(
        [{"code": k, "count": v} for k, v in signal_freq.items()],
        key=lambda x: x["count"],
        reverse=True,
    )[:5]

    return {
        "data": {
            "total_students": len(students),
            "at_risk_count": at_risk,
            "band_breakdown": band_counts,
            "top_signals": top_signals,
        },
        "meta": _meta(request),
    }


@router.get("/reports/dropout/students")
async def dropout_students(
    request: Request,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    band: Optional[str] = Query(None),
    sort: str = Query("risk_score"),
    order: str = Query("desc"),
    db: Session = Depends(get_db),
    school_id: uuid.UUID = Depends(get_school_id),
):
    """Paginated list of students with dropout risk scores."""
    token = _extract_token(request)
    today = date.today()

    all_students = list_active_students(db, school_id)
    enriched = []
    for student in all_students:
        sid_str = student.get("id")
        if not sid_str:
            continue
        try:
            sid_uuid = uuid.UUID(sid_str)
        except (TypeError, ValueError):
            continue
        risk = await compute_student_risk(db, school_id, sid_uuid, token, today)
        enriched.append({
            "student_id": sid_str,
            "student_code": student.get("student_code", ""),
            "first_name": student.get("first_name", ""),
            "last_name": student.get("last_name", ""),
            "risk_score": risk["risk_score"],
            "risk_band": risk["risk_band"],
            "signal_count": len(risk["signals"]),
            "top_signal": max(risk["signals"], key=lambda s: s["points"])["label"]
            if risk["signals"]
            else None,
        })

    if band:
        band_upper = band.upper()
        enriched = [s for s in enriched if s["risk_band"] == band_upper]

    reverse = order.lower() == "desc"
    enriched.sort(key=lambda s: s.get(sort, 0), reverse=reverse)

    total_filtered = len(enriched)
    start = (page - 1) * page_size
    page_data = enriched[start:start + page_size]

    meta = _meta(request)
    meta.update({
        "page": page,
        "page_size": page_size,
        "total": total_filtered,
        "has_next": (page * page_size) < total_filtered,
    })
    return {"data": page_data, "meta": meta}


@router.get("/reports/dropout/student/{student_id}")
async def dropout_student_detail(
    student_id: uuid.UUID,
    request: Request,
    db: Session = Depends(get_db),
    school_id: uuid.UUID = Depends(get_school_id),
):
    """Full dropout risk breakdown for a single student."""
    token = _extract_token(request)
    today = date.today()

    risk = await compute_student_risk(db, school_id, student_id, token, today)

    return {
        "data": {
            "student_id": str(student_id),
            "risk_score": risk["risk_score"],
            "risk_band": risk["risk_band"],
            "signals": risk["signals"],
            "computed_at": datetime.now(timezone.utc).isoformat(),
            "lookback_days": LOOKBACK_DAYS,
        },
        "meta": _meta(request),
    }


# ───────────────────── Exports (Excel / PDF) ───────────────────────


@router.get("/reports/export/attendance")
def export_attendance_excel(
    request: Request,
    from_date: date = Query(..., alias="from"),
    to_date: date = Query(..., alias="to"),
    rdb: Session = Depends(get_reporting_db),
    current_user: dict = Depends(get_current_user),
    school_id: uuid.UUID = Depends(get_school_id),
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return {"error": {"code": "DEPENDENCY_MISSING", "message": "openpyxl not installed"}}

    svc = ReportingQueryService(rdb)
    data = svc.get_attendance_trend(school_id, from_date, to_date)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Trend"
    ws.append(["Date", "Present", "Absent", "Late", "Total", "Rate (%)"])

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for row in data:
        ws.append([
            row.get("date", ""),
            row.get("present", 0),
            row.get("absent", 0),
            row.get("late", 0),
            row.get("total", 0),
            round(row.get("rate", 0), 1),
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col) + 2
        ws.column_dimensions[col[0].column_letter].width = max_len

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"attendance_{from_date}_{to_date}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/reports/export/financial")
def export_financial_excel(
    request: Request,
    year_id: Optional[uuid.UUID] = Query(None),
    rdb: Session = Depends(get_reporting_db),
    current_user: dict = Depends(get_current_user),
    school_id: uuid.UUID = Depends(get_school_id),
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return {"error": {"code": "DEPENDENCY_MISSING", "message": "openpyxl not installed"}}

    svc = ReportingQueryService(rdb)
    data = svc.get_financial_summary(school_id, year_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financial Summary"
    ws.append(["Academic Year ID", "Total Invoiced", "Total Paid", "Total Outstanding"])

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    if isinstance(data, dict):
        data = [data]

    for row in data:
        ws.append([
            str(row.get("academic_year_id", "")),
            row.get("total_invoiced", 0),
            row.get("total_paid", 0),
            row.get("total_outstanding", 0),
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=financial_summary.xlsx"},
    )


@router.get("/reports/export/dropout")
async def export_dropout_excel(
    request: Request,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
    school_id: uuid.UUID = Depends(get_school_id),
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return {"error": {"code": "DEPENDENCY_MISSING", "message": "openpyxl not installed"}}

    token = _extract_token(request)
    today = date.today()

    students = list_active_students(db, school_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dropout Risk"
    ws.append([
        "Student Code", "First Name", "Last Name",
        "Risk Score", "Risk Band", "Signals", "Top Signal",
    ])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="C62828", end_color="C62828", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for student in students:
        sid_str = student.get("id")
        if not sid_str:
            continue
        try:
            sid_uuid = uuid.UUID(sid_str)
        except (TypeError, ValueError):
            continue
        risk = await compute_student_risk(db, school_id, sid_uuid, token, today)
        top_sig = (
            max(risk["signals"], key=lambda s: s["points"])["label"]
            if risk["signals"]
            else ""
        )
        ws.append([
            student.get("student_code", ""),
            student.get("first_name", ""),
            student.get("last_name", ""),
            risk["risk_score"],
            risk["risk_band"],
            len(risk["signals"]),
            top_sig,
        ])

    risk_colors = {"LOW": "4CAF50", "MEDIUM": "FF9800", "HIGH": "FF5722", "CRITICAL": "B71C1C"}
    for row_idx in range(2, ws.max_row + 1):
        band_cell = ws.cell(row=row_idx, column=5)
        color = risk_colors.get(str(band_cell.value), "FFFFFF")
        band_cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        band_cell.font = Font(color="FFFFFF", bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=dropout_risk.xlsx"},
    )


@router.get("/reports/export/report-card/{student_id}")
async def export_report_card_pdf(
    student_id: uuid.UUID,
    request: Request,
    term_id: Optional[uuid.UUID] = Query(None),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
    school_id: uuid.UUID = Depends(get_school_id),
):
    """PDF report card. Student + marks both come from academics in-process now."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError:
        return {"error": {"code": "DEPENDENCY_MISSING", "message": "reportlab not installed"}}

    # Pull student + per-subject mark roll-up directly via in-process services
    # — the old HTTP hops into student-service and assessment-service are gone.
    from app.services.assessment_service import AssessmentService
    from app.services.student_service import StudentService

    s_svc = StudentService(db)
    student = s_svc.get_student(student_id, school_id) or {}

    try:
        a_svc = AssessmentService(db)
        marks = a_svc.student_marks(school_id, student_id, term_id=term_id)
    except (AttributeError, TypeError) as exc:
        # Backwards-compat: if AssessmentService doesn't expose this helper,
        # degrade gracefully to an empty marks list rather than 500.
        logger.warning(
            "reports.report_card.no_marks_helper",
            extra={"student_id": str(student_id), "err": str(exc)},
        )
        marks = []

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("<b>EDUZIM SCHOOL MANAGEMENT SYSTEM</b>", styles["Title"]))
    elements.append(Paragraph("STUDENT REPORT CARD", styles["Heading2"]))
    elements.append(Spacer(1, 10 * mm))

    name = f"{student.get('first_name', '')} {student.get('last_name', '')}".strip()
    elements.append(Paragraph(f"<b>Student:</b> {name}", styles["Normal"]))
    elements.append(Paragraph(
        f"<b>Admission No:</b> {student.get('student_code', student.get('admission_number', 'N/A'))}",
        styles["Normal"],
    ))
    elements.append(Paragraph(f"<b>Date:</b> {date.today().isoformat()}", styles["Normal"]))
    elements.append(Spacer(1, 8 * mm))

    table_data = [["Subject", "Assessments", "Average (%)", "Grade"]]
    for subj in marks or []:
        avg = subj.get("average_pct")
        grade = _grade(avg) if avg is not None else "N/A"
        table_data.append([
            (subj.get("subject_id", "") or "")[:12],
            str(subj.get("graded_count", 0)),
            f"{avg:.1f}" if avg is not None else "N/A",
            grade,
        ])

    if len(table_data) > 1:
        t = Table(table_data, colWidths=[50 * mm, 30 * mm, 35 * mm, 25 * mm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E7D32")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
        ]))
        elements.append(t)
    else:
        elements.append(Paragraph("No assessment data available for this period.", styles["Normal"]))

    elements.append(Spacer(1, 10 * mm))
    elements.append(Paragraph("_________________________", styles["Normal"]))
    elements.append(Paragraph("School Administrator", styles["Normal"]))

    doc.build(elements)
    buf.seek(0)

    filename = f"report_card_{student.get('last_name', 'student')}_{date.today()}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _grade(pct: float) -> str:
    """Zimbabwe O-Level grading scale."""
    if pct >= 80:
        return "A"
    elif pct >= 70:
        return "B"
    elif pct >= 60:
        return "C"
    elif pct >= 50:
        return "D"
    elif pct >= 40:
        return "E"
    else:
        return "U"
