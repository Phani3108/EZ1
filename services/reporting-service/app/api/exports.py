"""
Report Export Routes — PDF & Excel generation for reports and report cards.
"""
import io
import uuid
import logging
from datetime import date, datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.config import get_settings
from app.database import get_db
from app.dependencies import get_current_user, get_school_id
from app.services.reporting_service import ReportingService

logger = logging.getLogger(__name__)
router = APIRouter(tags=["Report Export"])

settings = get_settings()


def _meta(request: Request) -> dict:
    rid = getattr(request.state, "request_id", str(uuid.uuid4()))
    return {"request_id": rid, "timestamp": datetime.now(timezone.utc).isoformat()}


# ─── Excel Export ───

@router.get("/reports/export/attendance")
def export_attendance_excel(
    request: Request,
    from_date: date = Query(..., alias="from"),
    to_date: date = Query(..., alias="to"),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
    school_id: uuid.UUID = Depends(get_school_id),
):
    """Export attendance trend data as Excel (.xlsx)."""
    try:
        import openpyxl
    except ImportError:
        return {"error": {"code": "DEPENDENCY_MISSING", "message": "openpyxl not installed"}}

    svc = ReportingService(db)
    data = svc.get_attendance_trend(school_id, from_date, to_date)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Trend"

    # Header
    ws.append(["Date", "Present", "Absent", "Late", "Total", "Rate (%)"])

    # Style header
    from openpyxl.styles import Font, PatternFill
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    # Data rows
    for row in data:
        ws.append([
            row.get("date", ""),
            row.get("present", 0),
            row.get("absent", 0),
            row.get("late", 0),
            row.get("total", 0),
            round(row.get("rate", 0), 1),
        ])

    # Auto-width columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col) + 2
        ws.column_dimensions[col[0].column_letter].width = max_len

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"attendance_{from_date}_{to_date}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/reports/export/financial")
def export_financial_excel(
    request: Request,
    year_id: Optional[uuid.UUID] = Query(None),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
    school_id: uuid.UUID = Depends(get_school_id),
):
    """Export financial summary as Excel (.xlsx)."""
    try:
        import openpyxl
    except ImportError:
        return {"error": {"code": "DEPENDENCY_MISSING", "message": "openpyxl not installed"}}

    svc = ReportingService(db)
    data = svc.get_financial_summary(school_id, year_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financial Summary"

    ws.append(["Academic Year ID", "Total Invoiced", "Total Paid", "Total Outstanding"])

    from openpyxl.styles import Font, PatternFill, numbers
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    if isinstance(data, dict):
        data = [data]

    for row in data:
        ws.append([
            str(row.get("academic_year_id", "")),
            row.get("total_invoiced", 0),
            row.get("total_paid", 0),
            row.get("total_outstanding", 0),
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=financial_summary.xlsx"},
    )


@router.get("/reports/export/dropout")
async def export_dropout_excel(
    request: Request,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
    school_id: uuid.UUID = Depends(get_school_id),
):
    """Export dropout risk data as Excel (.xlsx)."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return {"error": {"code": "DEPENDENCY_MISSING", "message": "openpyxl not installed"}}

    import httpx
    token = request.headers.get("authorization", "").replace("Bearer ", "")
    today = date.today()
    sid = str(school_id)

    # Fetch students
    from app.api.routes import _fetch_students, _compute_student_risk
    students, _ = await _fetch_students(sid, token, page=1, page_size=500)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dropout Risk"

    ws.append(["Student Code", "First Name", "Last Name", "Risk Score", "Risk Band", "Signals", "Top Signal"])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="C62828", end_color="C62828", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for student in students:
        student_id = student.get("id")
        if not student_id:
            continue
        risk = await _compute_student_risk(sid, student_id, token, today)
        top_sig = max(risk["signals"], key=lambda s: s["points"])["label"] if risk["signals"] else ""
        ws.append([
            student.get("student_code", ""),
            student.get("first_name", ""),
            student.get("last_name", ""),
            risk["risk_score"],
            risk["risk_band"],
            len(risk["signals"]),
            top_sig,
        ])

    # Color-code risk bands
    risk_colors = {"LOW": "4CAF50", "MEDIUM": "FF9800", "HIGH": "FF5722", "CRITICAL": "B71C1C"}
    for row_idx in range(2, ws.max_row + 1):
        band_cell = ws.cell(row=row_idx, column=5)
        color = risk_colors.get(str(band_cell.value), "FFFFFF")
        band_cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        band_cell.font = Font(color="FFFFFF", bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=dropout_risk.xlsx"},
    )


# ─── PDF Report Card ───

@router.get("/reports/export/report-card/{student_id}")
async def export_report_card_pdf(
    student_id: uuid.UUID,
    request: Request,
    term_id: Optional[uuid.UUID] = Query(None),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
    school_id: uuid.UUID = Depends(get_school_id),
):
    """Generate a PDF report card for a student."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError:
        return {"error": {"code": "DEPENDENCY_MISSING", "message": "reportlab not installed"}}

    import httpx
    token = request.headers.get("authorization", "").replace("Bearer ", "")

    # Fetch student info and marks
    async with httpx.AsyncClient(timeout=10.0) as client:
        # Student details
        student_resp = await client.get(
            f"{settings.STUDENT_SERVICE_URL}/api/v1/students/{student_id}",
            headers={"Authorization": f"Bearer {token}"},
        )
        student = student_resp.json().get("data", {}) if student_resp.status_code == 200 else {}

        # Student marks
        params = {}
        if term_id:
            params["term_id"] = str(term_id)
        marks_resp = await client.get(
            f"{settings.ASSESSMENT_SERVICE_URL}/api/v1/assessments/students/{student_id}/marks",
            params=params,
            headers={"Authorization": f"Bearer {token}"},
        )
        marks = marks_resp.json().get("data", []) if marks_resp.status_code == 200 else []

    # Build PDF
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()
    elements = []

    # School header
    elements.append(Paragraph("<b>EDUZIM SCHOOL MANAGEMENT SYSTEM</b>", styles["Title"]))
    elements.append(Paragraph("STUDENT REPORT CARD", styles["Heading2"]))
    elements.append(Spacer(1, 10 * mm))

    # Student info
    name = f"{student.get('first_name', '')} {student.get('last_name', '')}"
    elements.append(Paragraph(f"<b>Student:</b> {name}", styles["Normal"]))
    elements.append(Paragraph(f"<b>Admission No:</b> {student.get('student_code', student.get('admission_number', 'N/A'))}", styles["Normal"]))
    elements.append(Paragraph(f"<b>Date:</b> {date.today().isoformat()}", styles["Normal"]))
    elements.append(Spacer(1, 8 * mm))

    # Marks table
    table_data = [["Subject", "Assessments", "Average (%)", "Grade"]]
    for subj in marks:
        avg = subj.get("average_pct")
        grade = _grade(avg) if avg is not None else "N/A"
        table_data.append([
            subj.get("subject_id", "")[:12],
            str(subj.get("graded_count", 0)),
            f"{avg:.1f}" if avg is not None else "N/A",
            grade,
        ])

    if len(table_data) > 1:
        t = Table(table_data, colWidths=[50 * mm, 30 * mm, 35 * mm, 25 * mm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E7D32")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
        ]))
        elements.append(t)
    else:
        elements.append(Paragraph("No assessment data available for this period.", styles["Normal"]))

    elements.append(Spacer(1, 10 * mm))
    elements.append(Paragraph("_________________________", styles["Normal"]))
    elements.append(Paragraph("School Administrator", styles["Normal"]))

    doc.build(elements)
    buf.seek(0)

    filename = f"report_card_{student.get('last_name', 'student')}_{date.today()}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _grade(pct: float) -> str:
    """Zimbabwe O-Level grading scale."""
    if pct >= 80:
        return "A"
    elif pct >= 70:
        return "B"
    elif pct >= 60:
        return "C"
    elif pct >= 50:
        return "D"
    elif pct >= 40:
        return "E"
    else:
        return "U"
